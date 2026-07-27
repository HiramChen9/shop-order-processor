#!/usr/bin/env python3
"""
E-commerce order Excel processor for Shopee, Lazada, TikTok Shop.
"""
import argparse, os, re, sys
try:
    import openpyxl
    from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
    import openpyxl.utils
except ImportError:
    print("Error: openpyxl required"); sys.exit(1)

SKU_RULES = [
    (re.compile(r'^CZTWC.+-3060\*20$'), 26.4, 0),
    (re.compile(r'^CZTWC.+-3060\*10$'), 13.2, 0),
    (re.compile(r'^BKCZT-.+-3060\*20$'), 13.6, 1.07524864),
    (re.compile(r'^BKCZT-.+-3060\*10$'), 6.8, 0.53762432),
    (re.compile(r'^CZT-.+-4090-1\.0\*5$'), 8.25, 2.0017452),
    (re.compile(r'^CZT-.+-4090\*5$'), 9.2, 4.0034904),
    (re.compile(r'^CZT-.+-4090$'), 1.84, 0.80069808),
    (re.compile(r'^CZT-.+-3060-1\.0\*20$'), 16.48, 3.6368704),
    (re.compile(r'^CZT-.+-3060-1\.0\*30$'), 24.72, 5.4553056),
    (re.compile(r'^CZT-.+-3060\*20$'), 19, 8.0643648),
    (re.compile(r'^CZT-.+-3060\*30$'), 28.5, 12.0965472),
    (re.compile(r'^CZT-.+-3060-1\.0\*10$'), 8.24, 1.8184352),
    (re.compile(r'^CZT-.+-3060\*10$'), 9.5, 4.0321824),
    (re.compile(r'^CZTZHZ-.+$'), 27.52, 12.3683242),
    (re.compile(r'^XZ-.+-3$'), 3.06, 0.8301552),
    (re.compile(r'^(?:N|DE|MA).+-3$'), 2.8, 0.57843072),
    (re.compile(r'^PPD.+-8$'), 27.21, 0),
    (re.compile(r'^PPD.+-10$'), 31.6, 0),
    (re.compile(r'^YXX-.+-45-5$'), 2.84, 0.51131535),
    (re.compile(r'^(?:AN|DE|MA).+-5$'), 4.45, 0.7230384),
    (re.compile(r'^CZTJZH-.+-120-2\.8$'), 19.28, 10.549889),
    (re.compile(r'^MBT.+-10-5$'), 2.83, 1.42751555555556),
]


def parse_idr(val):
    if val is None: return None
    s = str(val).strip()
    if not s or s == "0": return 0
    s = s.replace(".", "")
    try: return int(s)
    except ValueError:
        try: return float(s)
        except ValueError: return val


def clean_sku(sku):
    if not sku: return ""
    s = str(sku).strip()
    return re.sub(r'-[A-Z]+$', '', s)


def trunc_date(val):
    if val is None: return None
    s = str(val).strip()
    if len(s) >= 10 and s[4] == "-" and s[7] == "-":
        return s[:10]
    m = re.match(r'^(\d{2})/(\d{2})/(\d{4})\s', s)
    if m: return f"{m.group(3)}-{m.group(2)}-{m.group(1)}"
    m = re.match(r'^(\d{1,2})\s+(\w+)\s+(\d{4})\s', s)
    if m:
        month_map = {"Jan":"01","Feb":"02","Mar":"03","Apr":"04","May":"05","Jun":"06",
                     "Jul":"07","Aug":"08","Sep":"09","Oct":"10","Nov":"11","Dec":"12"}
        mon = month_map.get(m.group(2))
        if mon: return f"{m.group(3)}-{mon}-{int(m.group(1)):02d}"
    return s


def detect_platform(hd):
    """Auto-detect platform from header dict."""
    keys = list(hd.keys())
    kset = set(k.lower() for k in keys)
    # Shopee: Indonesian columns first (no. pesanan is unique)
    if 'no. pesanan' in kset or 'nomor referensi sku' in kset or 'subtotal pesanan' in kset:
        return 'shopee'
    if 'status pesanan' in kset or 'voucher ditanggung penjual' in kset:
        return 'shopee'
    # Lazada: sellerSku + createTime/orderNumber/paidPrice
    if 'sellerSku' in hd or 'seller sku' in kset:
        if 'createTime' in hd or 'create time' in kset:
            return 'lazada'
        if 'orderNumber' in hd or 'order number' in kset:
            return 'lazada'
        if 'paidPrice' in hd or 'paid price' in kset:
            return 'lazada'
    # TikTok: specific columns
    if 'order id' in kset or 'order status' in kset:
        if 'normal or pre-order' in kset:
            return 'tiktok'
        if 'sku platform discount' in kset or 'payment platform discount' in kset:
            return 'tiktok'
    # Fallbacks
    if 'sellerSku' in hd: return 'lazada'
    if 'Order ID' in hd and 'Order Status' in hd: return 'tiktok'
    return None


def match_cost_shipping(sku):
    for p, c, s in SKU_RULES:
        if p.match(sku): return c, s
    return None, None


def is_black_bg(cell):
    f = cell.fill
    return f and f.start_color and f.start_color.rgb == "00000000"


def style_header(ws, headers):
    fl = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    fn = Font(bold=True, size=11, color="FFFFFF")
    bd = Border(left=Side(style="thin"),right=Side(style="thin"),top=Side(style="thin"),bottom=Side(style="thin"))
    for ci, h in enumerate(headers, 1):
        c = ws.cell(1, ci, h); c.font = fn; c.fill = fl; c.alignment = Alignment(horizontal="center"); c.border = bd


def auto_width(ws, headers):
    for ci, h in enumerate(headers, 1):
        ml = max(len(h.encode("utf-8")), 8)
        for ri in range(2, ws.max_row + 1):
            if ws.cell(ri, ci).value is not None:
                ml = max(ml, len(str(ws.cell(ri, ci).value).encode("utf-8")))
        ws.column_dimensions[openpyxl.utils.get_column_letter(ci)].width = min(ml // 2 + 6, 55)


def write_data_sheet(ws, rows, headers):
    style_header(ws, headers)
    bd = Border(left=Side(style="thin"),right=Side(style="thin"),top=Side(style="thin"),bottom=Side(style="thin"))
    cidx = next((i for i, h in enumerate(headers) if h in ("\u6210\u672c", "cost")), None)
    sidx = next((i for i, h in enumerate(headers) if h in ("\u8fd0\u8d39", "shipping")), None)
    ctx = next((i for i, h in enumerate(headers) if h == "\u6210\u672c\u5408\u8ba1"), None)
    stx = next((i for i, h in enumerate(headers) if h == "\u8fd0\u8d39\u5408\u8ba1"), None)
    qn = next((i for i, h in enumerate(headers) if h in ("\u4ef6\u6570", "Quantity")), None)
    mn = next((i for i, h in enumerate(headers) if h in ("\u8ba2\u5355\u5c0f\u8ba1", "paidPrice", "SKU Subtotal After Discount")), None)
    for ri, rd in enumerate(rows, 2):
        for ci, val in enumerate(rd, 1):
            c = ws.cell(ri, ci, val); c.border = bd
            if ci in (cidx, sidx, ctx, stx) and val is not None:
                c.alignment = Alignment(horizontal="right"); c.number_format = "#,##0.00"
            elif ci == mn:
                c.alignment = Alignment(horizontal="right")
                if val is not None: c.number_format = "#,##0"
            elif ci == qn: c.alignment = Alignment(horizontal="right")
            else: c.alignment = Alignment(horizontal="left")
    auto_width(ws, headers)


def add_summary(wb, rows, headers, date_idx, order_idx, qty_idx, cost_total_idx, ship_total_idx, sales_fn=None):
    ws = wb.create_sheet(title="\u6bcf\u65e5\u7edf\u8ba1")
    sum_headers = ["\u65e5\u671f", "\u8ba2\u5355\u603b\u6570", "\u4ef6\u6570\u603b\u6570"]
    sub_idx = next((i for i, h in enumerate(headers) if h in ("\u8ba2\u5355\u5c0f\u8ba1", "paidPrice", "SKU Subtotal After Discount")), None)
    if sub_idx is not None: sum_headers.append("\u91d1\u989d\u5c0f\u8ba1")
    vou_idx = next((i for i, h in enumerate(headers) if h == "\u5356\u5bb6\u627f\u62c5\u7684\u4f18\u60e0\u5238"), None)
    if vou_idx is not None: sum_headers.append("\u4f18\u60e0\u5238")
    if cost_total_idx is not None: sum_headers.append("\u6210\u672c\u603b\u989d")
    if ship_total_idx is not None: sum_headers.append("\u8fd0\u8d39\u603b\u989d")
    if sales_fn is not None: sum_headers.append("\u5b9e\u9645\u9500\u552e\u91d1\u989d")
    style_header(ws, sum_headers)
    bd = Border(left=Side(style="thin"),right=Side(style="thin"),top=Side(style="thin"),bottom=Side(style="thin"))
    daily = {}
    for rd in rows:
        d = str(rd[date_idx]) if rd[date_idx] else "unknown"
        if d not in daily:
            daily[d] = {"orders": set(), "qty": 0, "sub": 0, "vou": 0, "ct": 0.0, "st": 0.0, "sales": 0.0}
        oid = str(rd[order_idx]) if rd[order_idx] else ""
        daily[d]["orders"].add(oid)
        if qty_idx is not None and rd[qty_idx] is not None:
            daily[d]["qty"] += int(rd[qty_idx]) if isinstance(rd[qty_idx], int) else 0
        if sub_idx is not None and rd[sub_idx] is not None:
            daily[d]["sub"] += float(rd[sub_idx]) if isinstance(rd[sub_idx], (int, float)) else 0
        if vou_idx is not None and rd[vou_idx] is not None:
            daily[d]["vou"] += float(rd[vou_idx]) if isinstance(rd[vou_idx], (int, float)) else 0
        if cost_total_idx is not None and rd[cost_total_idx] is not None:
            daily[d]["ct"] += float(rd[cost_total_idx])
        if ship_total_idx is not None and rd[ship_total_idx] is not None:
            daily[d]["st"] += float(rd[ship_total_idx])
        if sales_fn is not None:
            daily[d]["sales"] += float(sales_fn(rd))
    row_idx = 2
    for date in sorted(daily.keys()):
        d = daily[date]
        vals = [date, len(d["orders"]), d["qty"]]
        if sub_idx is not None: vals.append(d["sub"])
        if vou_idx is not None: vals.append(d["vou"])
        if cost_total_idx is not None: vals.append(d["ct"])
        if ship_total_idx is not None: vals.append(d["st"])
        if sales_fn is not None: vals.append(d["sales"])
        for ci, val in enumerate(vals, 1):
            c = ws.cell(row_idx, ci, val); c.border = bd
            if ci == 1: c.alignment = Alignment(horizontal="left")
            else:
                c.alignment = Alignment(horizontal="right")
                if ci >= 3: c.number_format = "#,##0"
        row_idx += 1
    tf = Font(bold=True, size=11)
    tv = ["\u5408\u8ba1", sum(len(v["orders"]) for v in daily.values()), sum(v["qty"] for v in daily.values())]
    if sub_idx is not None: tv.append(sum(v["sub"] for v in daily.values()))
    if vou_idx is not None: tv.append(sum(v["vou"] for v in daily.values()))
    if cost_total_idx is not None: tv.append(sum(v["ct"] for v in daily.values()))
    if ship_total_idx is not None: tv.append(sum(v["st"] for v in daily.values()))
    if sales_fn is not None: tv.append(sum(v["sales"] for v in daily.values()))
    for ci, val in enumerate(tv, 1):
        c = ws.cell(row_idx, ci, val); c.border = bd; c.font = tf
        if ci == 1: c.alignment = Alignment(horizontal="left")
        else:
            c.alignment = Alignment(horizontal="right")
            if ci >= 3: c.number_format = "#,##0"
    auto_width(ws, sum_headers)


def process_shopee(in_path, out_path):
    wb = openpyxl.load_workbook(in_path); ws = wb.active
    excl = {"Batal", "Belum Bayar", "Pembatalan diajukan"}
    headers = ["订单号", "订单创建时间", "参考sku编码",
               "成本", "运费", "件数", "成本合计", "运费合计",
               "订单小计", "卖家承担的优惠券"]
    rows = []
    for r in range(2, ws.max_row + 1):
        if ws.cell(r, 2).value in excl: continue
        black = is_black_bg(ws.cell(r, 1))
        raw = ws.cell(r, 15).value or ""
        cost, ship = match_cost_shipping(clean_sku(raw))
        try: qty = int(str(ws.cell(r, 19).value or "0").strip())
        except: qty = 0
        sub = parse_idr(ws.cell(r, 21).value)
        vou = parse_idr(ws.cell(r, 28).value)
        if black: vou = 0
        ct = round(cost * qty, 2) if cost is not None else None
        st = round(ship * qty, 2) if ship is not None else None
        date_val = trunc_date(ws.cell(r, 10).value)
        rows.append([ws.cell(r, 1).value, date_val, raw, cost, ship, qty, ct, st, sub, vou])
    wb_out = openpyxl.Workbook()
    ws_out = wb_out.active; ws_out.title = "筛选后"
    write_data_sheet(ws_out, rows, headers)
    add_summary(wb_out, rows, headers, date_idx=1, order_idx=0, qty_idx=5,
                cost_total_idx=6, ship_total_idx=7,
                sales_fn=lambda r: (r[8] or 0) - (r[9] or 0))
    wb_out.save(out_path); return len(rows)


def process_lazada(in_path, out_path):
    wb = openpyxl.load_workbook(in_path); ws = wb.active
    hd = {}
    for c in range(1, ws.max_column + 1):
        v = ws.cell(1, c).value
        if v: hd[v.strip()] = c
    for k in ("sellerSku", "createTime", "orderNumber", "paidPrice", "status"):
        if k not in hd: print("Error: column " + k + " not found"); sys.exit(1)
    ci = {k: hd[k] for k in ("sellerSku", "createTime", "orderNumber", "paidPrice", "status")}
    excl = {"canceled", "cancelled", "unpaid", "batal", "belum bayar"}
    headers = ["sellerSku", "createTime", "orderNumber", "paidPrice", "status",
               "成本", "运费", "件数", "成本合计", "运费合计"]
    raw_rows = []
    for r in range(2, ws.max_row + 1):
        st = str(ws.cell(r, ci["status"]).value or "").strip().lower()
        if any(e in st for e in excl): continue
        raw_sku = ws.cell(r, ci["sellerSku"]).value or ""
        cost, ship = match_cost_shipping(clean_sku(raw_sku))
        date_val = trunc_date(ws.cell(r, ci["createTime"]).value)
        raw_rows.append({"sku": raw_sku, "date": date_val,
            "order": ws.cell(r, ci["orderNumber"]).value,
            "price": parse_idr(ws.cell(r, ci["paidPrice"]).value),
            "status": ws.cell(r, ci["status"]).value,
            "cost": cost, "ship": ship})
    rows = []
    for rr in raw_rows:
        qty = 1
        ct = rr["cost"]
        ste = rr["ship"]
        rows.append([rr["sku"], rr["date"], rr["order"], rr["price"], rr["status"],
                     rr["cost"], rr["ship"], qty, ct, ste])
    wb_out = openpyxl.Workbook()
    ws_out = wb_out.active; ws_out.title = "筛选后"
    write_data_sheet(ws_out, rows, headers)
    add_summary(wb_out, rows, headers, date_idx=1, order_idx=2, qty_idx=7,
                cost_total_idx=8, ship_total_idx=9,
                sales_fn=lambda r: r[3] or 0)
    wb_out.save(out_path); return len(rows)


def process_tiktok(in_path, out_path):
    wb = openpyxl.load_workbook(in_path); ws = wb.active
    hd = {}
    for c in range(1, ws.max_column + 1):
        v = ws.cell(1, c).value
        if v: hd[v.strip()] = c
    req = ["Order ID", "Order Status", "Normal or Pre-order", "Seller SKU", "Quantity",
           "SKU Platform Discount", "SKU Subtotal After Discount", "Payment platform discount", "Created Time"]
    for k in req:
        if k not in hd: print("Error: column " + k + " not found"); sys.exit(1)
    ci = {k: hd[k] for k in req}
    excl = {"canceled", "cancelled", "unpaid", "batal", "belum bayar"}
    headers = ["Order ID", "Order Status", "Normal or Pre-order", "Seller SKU",
               "成本", "运费", "Quantity", "成本合计", "运费合计",
               "SKU Platform Discount", "SKU Subtotal After Discount", "Payment platform discount", "Created Time"]
    rows = []
    for r in range(2, ws.max_row + 1):
        st = str(ws.cell(r, ci["Order Status"]).value or "").strip().lower()
        if any(e in st for e in excl): continue
        pre = ws.cell(r, ci["Normal or Pre-order"]).value
        if pre is None or str(pre).strip() == "": continue
        raw = ws.cell(r, ci["Seller SKU"]).value or ""
        cost, ship = match_cost_shipping(clean_sku(raw))
        try: qty = int(str(ws.cell(r, ci["Quantity"]).value or "0").strip())
        except: qty = 0
        ct = round(cost * qty, 2) if cost is not None else None
        ste = round(ship * qty, 2) if ship is not None else None
        date_val = trunc_date(ws.cell(r, ci["Created Time"]).value)
        rows.append([ws.cell(r, ci["Order ID"]).value, ws.cell(r, ci["Order Status"]).value, pre, raw,
                     cost, ship, qty, ct, ste,
                     parse_idr(ws.cell(r, ci["SKU Platform Discount"]).value),
                     parse_idr(ws.cell(r, ci["SKU Subtotal After Discount"]).value),
                     parse_idr(ws.cell(r, ci["Payment platform discount"]).value),
                     date_val])
    wb_out = openpyxl.Workbook()
    ws_out = wb_out.active; ws_out.title = "筛选后"
    write_data_sheet(ws_out, rows, headers)
    add_summary(wb_out, rows, headers, date_idx=12, order_idx=0, qty_idx=6,
                cost_total_idx=7, ship_total_idx=8,
                sales_fn=lambda r: (r[9] or 0) + (r[10] or 0) + (r[11] or 0))
    wb_out.save(out_path); return len(rows)


def main():
    p = argparse.ArgumentParser(description="Process e-commerce order Excel files.")
    p.add_argument("platform", nargs="?", default="auto",
                    choices=["shopee", "lazada", "tiktok", "auto"],
                    help="Platform type (auto = detect from file headers)")
    p.add_argument("input", help="Input Excel file path")
    p.add_argument("output", nargs="?", default=None, help="Output path")
    a = p.parse_args()
    if not os.path.exists(a.input):
        print("Error: " + a.input + " not found"); sys.exit(1)
    plat = a.platform
    if plat == "auto":
        wb = openpyxl.load_workbook(a.input)
        ws = wb.active
        hd = {}
        for c in range(1, ws.max_column + 1):
            v = ws.cell(1, c).value
            if v: hd[v.strip()] = c
        wb.close()
        plat = detect_platform(hd)
        if not plat:
            print("Error: cannot detect platform. Headers found:")
            for k in hd: print("  " + repr(k))
            print("Specify platform: shopee|lazada|tiktok")
            sys.exit(1)
        print("Detected platform: " + plat)
    base, ext = os.path.splitext(a.input)
    out = a.output or (base + "_筛选后" + ext)
    fn = {"shopee": process_shopee, "lazada": process_lazada, "tiktok": process_tiktok}[plat]
    n = fn(a.input, out)
    print("Done: " + str(n) + " rows written -> " + out)

if __name__ == "__main__": main()
